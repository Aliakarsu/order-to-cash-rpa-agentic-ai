"""
common.py - Shared helpers: logging, simulated e-mail sending, Excel access.

E-mail note: in production this module would call Outlook via win32com or an
SMTP server (as covered in Lecture 8 / W8 examples). For a portable,
gradeable demo it writes RFC-822 style .eml files into data/outbox/ instead,
which can be opened with any mail client. Switch SMTP_ENABLED in config.py
to send real mail.
"""
import logging
import smtplib
from datetime import datetime
from email.message import EmailMessage

import openpyxl

import config


def get_logger(name: str) -> logging.Logger:
    config.ensure_dirs()
    logger = logging.getLogger(name)
    if logger.handlers:
        return logger
    logger.setLevel(logging.INFO)
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s %(name)s: %(message)s")
    fh = logging.FileHandler(config.LOG_DIR / "run.log", encoding="utf-8")
    ch = logging.StreamHandler()
    for h in (fh, ch):
        h.setFormatter(fmt)
        logger.addHandler(h)
    return logger


def send_email(to_addr: str, subject: str, body: str,
               attachment_path=None, logger=None) -> str:
    """Send an e-mail, or simulate it by writing a .eml file to the outbox."""
    msg = EmailMessage()
    msg["From"] = "orders@wholesaleco.ie"
    msg["To"] = to_addr
    msg["Subject"] = subject
    msg["Date"] = datetime.now().strftime("%a, %d %b %Y %H:%M:%S +0100")
    msg.set_content(body)
    if attachment_path is not None:
        with open(attachment_path, "rb") as f:
            data = f.read()
        msg.add_attachment(data, maintype="application",
                           subtype="octet-stream",
                           filename=attachment_path.name)

    if config.SMTP_ENABLED:
        with smtplib.SMTP(config.SMTP_HOST, config.SMTP_PORT) as s:
            s.starttls()
            s.login(config.SMTP_USER, config.SMTP_PASSWORD)
            s.send_message(msg)
        target = f"SMTP:{to_addr}"
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        path = config.OUTBOX_DIR / f"{stamp}_{subject[:30].replace(' ', '_')}.eml"
        with open(path, "wb") as f:
            f.write(bytes(msg))
        target = str(path.name)

    if logger:
        logger.info("E-mail '%s' -> %s (%s)", subject, to_addr, target)
    return target


# ---------------- Excel helpers (openpyxl) ----------------

def read_table(xlsx_path, sheet=None):
    """Read a worksheet whose first row is a header; return list of dicts."""
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[sheet] if sheet else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def write_table(xlsx_path, header, rows, sheet_title="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header])
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, len(h) + 4)
    wb.save(xlsx_path)
